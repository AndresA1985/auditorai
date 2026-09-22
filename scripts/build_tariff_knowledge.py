"""Read ONLY the two approved worksheets; export deidentified precedents.

Run with bundled Python/openpyxl. No workbook is written. All clinical bodies
stay in memory; the output allowlist contains only semantic tags and tariffs.
"""
import argparse
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.tariff_semantics import VERSION, clinical_features, site_labels


def select_sheets(workbook):
    sheets = {}
    for target in ("hallazgos_base", "HONORARIOS"):
        names = [name for name in workbook.sheetnames if name.strip() == target]
        if len(names) != 1:
            raise ValueError("Hoja ausente o ambigua tras trim: " + target)
        sheets[target] = workbook[names[0]]
    return sheets


def parse_codes(value):
    if value is None or str(value).strip() == "": return [], False
    text = str(value).strip()
    if isinstance(value, (int, float)):
        text = str(int(value)) if int(value) == value else text
    # Parenthetical prose is not scanned for numbers. A trailing explanation
    # after the second alternative is recognized only in the source's format.
    alt = re.fullmatch(r"(\d{5,8})\s+O\s+(\d{5,8})\s*(?:\(?DEPENDIENDO.*)?", text, re.I)
    if alt: return [alt.group(1), alt.group(2)], True
    if not re.fullmatch(r"\d{5,8}(?:\s*[+\-]\s*\d{5,8})*", text):
        raise ValueError("Formato de códigos no reconocido")
    return re.split(r"\s*[+\-]\s*", text), False


def parse_percentages(value, number_format="General"):
    if value is None or str(value).strip() == "": return []
    if isinstance(value, bool): raise ValueError("Porcentaje booleano")
    if isinstance(value, (int,float)):
        amount = Decimal(str(value)) * (100 if "%" in number_format else 1)
        values = [amount]
    else:
        text = str(value).strip()
        if not re.fullmatch(r"\d+(?:[.,]\d+)?\s*%?(?:\s*[+\-]\s*\d+(?:[.,]\d+)?\s*%?)*", text):
            raise ValueError("Formato de porcentajes no reconocido")
        values = [Decimal(part.strip().replace("%", "").strip().replace(",","."))
                  for part in re.split(r"[+\-]", text)]
    if any(not value.is_finite() or not 0 <= value <= 100 for value in values):
        raise ValueError("Porcentaje fuera de rango")
    return [int(v) if v == int(v) else float(v) for v in values]


def parse_time(value):
    if value is None or str(value).strip() == "": return None
    if isinstance(value, bool): raise ValueError("Tiempo booleano")
    try: number = Decimal(str(value).strip())
    except InvalidOperation: raise ValueError("Tiempo no escalar") from None
    if not number.is_finite() or number <= 0: raise ValueError("Tiempo no válido")
    return int(number) if number == int(number) else float(number)


def directive_percentage(value):
    values={int(v) for v in re.findall(r"\b(\d+)\s*%",str(value or ""))}
    if len(values)!=1 or not 0<=next(iter(values))<=100:
        raise ValueError("Directriz sin porcentaje inequívoco")
    return next(iter(values))


def compile_workbook(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = select_sheets(workbook)
        base, fees = sheets["hallazgos_base"], sheets["HONORARIOS"]
        expected = {"Q":"procedimiento_sistema", "Y":"hallazgos_conclusion", "Z":"descripcion_estudio_013",
                    "AA":"codigo_auditor", "AB":"honorarios", "AC":"tiempo_anestesia", "AD":"nombre_procedimiento"}
        for column,label in expected.items():
            if str(base[column+"1"].value).strip() != label: raise ValueError("Cabecera inesperada: " + column)
        records = []
        columns={"Q":16,"Y":24,"Z":25,"AA":26,"AB":27,"AC":28,"AD":29,"AE":30,"AF":31}
        for row, source_row in enumerate(base.iter_rows(min_row=2),2):
            cells = {col:source_row[index] for col,index in columns.items()}
            issues = []
            codes, alternative = [], False
            percentages, anesthesia = [], None
            try: codes, alternative = parse_codes(cells["AA"].value)
            except ValueError: issues.append("codigos_no_interpretables")
            try: percentages = parse_percentages(cells["AB"].value, cells["AB"].number_format)
            except ValueError: issues.append("honorarios_no_interpretables")
            try: anesthesia = parse_time(cells["AC"].value)
            except ValueError: issues.append("anestesia_no_interpretable")
            if not codes: issues.append("sin_etiqueta")
            if len(percentages) != (1 if alternative else len(codes)):
                issues.append("cardinalidad_honorarios")
            if codes and anesthesia is None: issues.append("anestesia_ausente")
            clinical = clinical_features({"hallazgos_conclusion":cells["Y"].value,
                                          "descripcion_estudio_013":cells["Z"].value})
            if not clinical["actos"] or len(clinical["familias"]) != 1:
                issues.append("evidencia_semantica_insuficiente")
            if clinical["incertidumbres"]: issues.append("evidencia_semantica_incierta")
            # Notes are read, but neither medications nor supplies become codes.
            note = str(cells["AE"].value or "").lower()
            exclusive = alternative and "paquete" in note
            records.append({"fila":row,"codigos":codes,"alternativa":alternative,
                "porcentajes":percentages,"tiempo_anestesia":anesthesia,
                "caracteristicas":clinical,"incidencias":issues,
                "paquete_excluye_adicionales":exclusive,
                "directriz_sala_presente":bool(cells["AF"].value),
                "celdas":{"clinica":[f"Y{row}",f"Z{row}"],"contexto":[f"Q{row}",f"AD{row}"],
                          "codigos":f"AA{row}","porcentajes":f"AB{row}","anestesia":f"AC{row}","notas":f"AE{row}"}})
        groups=defaultdict(list)
        for record in records:
            if record["codigos"] and not record["alternativa"]: groups[tuple(record["codigos"])].append(record)
        for group in groups.values():
            distributions = {tuple(r["porcentajes"]) for r in group}
            times = {r["tiempo_anestesia"] for r in group}
            for record in group:
                if len(distributions)>1: record["incidencias"].append("honorarios_contradictorios")
                if len(times)>1: record["incidencias"].append("anestesia_no_uniforme")
        biopsies=[]
        for row in range(14,18):
            text=str(fees[f"B{row}"].value or "")
            match=re.match(r"\s*(\d{6})\s*\(",text)
            if not match: raise ValueError("Regla de biopsia inesperada")
            biopsies.append({"codigo":match.group(1),"sitios":sorted(site_labels(text)),"celda":f"B{row}"})
        additional=re.findall(r"\b\d{5,8}\b",str(fees["B10"].value))
        room_codes=re.findall(r"\b\d{6}\b",str(fees["A3"].value))
        if len(room_codes)!=4: raise ValueError("Directriz sala inesperada")
        rules={"adicionales_100":{"codigos":additional,"porcentaje":directive_percentage(fees["A9"].value),"celdas":["A9","B10"]},
               "principal":{"porcentaje":directive_percentage(fees["A5"].value),"celdas":["A5"]},
               "adicionales_50":{"porcentaje":directive_percentage(fees["A7"].value),"codigos":[],"celdas":["A7"]},
               "cpre_habitual":{"codigo":str(fees["B12"].value),"celdas":["A11","B12"]},
               "biopsias":biopsies,"cantidad_biopsias":{"base":"informes_013b","celdas":["A13"]},
               "sala":{"crm_segundo":room_codes[:2],"gastro_tercero":room_codes[2:],
                       "porcentaje":directive_percentage(fees["A3"].value),"solo_abiertos":True,"celdas":["A2","A3"]}}
        return {"schema_version":"1","semantics_version":VERSION,
                "fuente_sha256":hashlib.sha256(Path(path).read_bytes()).hexdigest(),
                "hojas":{"precedentes":base.title,"directrices":fees.title},
                "registros":records,"reglas":rules,
                "resumen":{"filas":len(records),"etiquetadas":sum(bool(r["codigos"]) for r in records),
                           "referencias_semanticas_utilizables":sum(bool(r["codigos"]) and not any(i.startswith("evidencia_semantica") for i in r["incidencias"]) for r in records),
                           "incidencias":dict(Counter(i for r in records for i in r["incidencias"]))}}
    finally: workbook.close()


if __name__ == "__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("workbook",type=Path)
    parser.add_argument("output",type=Path)
    args=parser.parse_args()
    result=compile_workbook(args.workbook)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8",newline="\n")
    print(json.dumps(result["resumen"]))
